import base64
import io

from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)

CALCULATE_DISCOUNT = """\
DECLARE total NUMBER DEFAULT 0;
DECLARE discount NUMBER DEFAULT 0;
SET total = price * quantity;
IF total > 100 THEN
    SET discount = total * 0.1;
ELSE
    SET discount = total * 0.05;
END IF;
SET total = total - discount;
"""


def _run_debug(params=None):
    """Drives the real /debug endpoint first, exactly like the frontend
    does before ever calling /debug/report -- so every test here exports
    a genuine, freshly-interpreted trace, never a hand-built fixture."""
    response = client.post(
        "/debug",
        json={"code": CALCULATE_DISCOUNT, "params": params or {"price": 20, "quantity": 6}, "name": "Calculate Discount"},
    )
    assert response.status_code == 200
    return response.json()["steps"]


def _tiny_png_data_url():
    from PIL import Image as PILImage

    buf = io.BytesIO()
    PILImage.new("RGB", (30, 15), color=(0, 128, 128)).save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def test_report_endpoint_rejects_unknown_format():
    steps = _run_debug()
    response = client.post(
        "/debug/report",
        json={"format": "exe", "code": CALCULATE_DISCOUNT, "steps": steps},
    )
    assert response.status_code == 400


def test_report_endpoint_rejects_empty_steps():
    response = client.post(
        "/debug/report",
        json={"format": "txt", "code": CALCULATE_DISCOUNT, "steps": []},
    )
    assert response.status_code == 400


def test_report_endpoint_txt_contains_real_run_data():
    steps = _run_debug(params={"price": 20, "quantity": 6})
    response = client.post(
        "/debug/report",
        json={
            "format": "txt",
            "code": CALCULATE_DISCOUNT,
            "params": {"price": 20, "quantity": 6},
            "name": "Calculate Discount",
            "steps": steps,
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/plain")
    assert "attachment" in response.headers["content-disposition"]
    assert ".txt" in response.headers["content-disposition"]

    text = response.text
    assert "Calculate Discount" in text
    assert "108.0" in text  # the real final total, same value /debug itself returned
    assert "lorem ipsum" not in text.lower()


def test_report_endpoint_docx_opens_and_contains_real_run_data():
    from docx import Document

    steps = _run_debug(params={"price": 20, "quantity": 6})
    response = client.post(
        "/debug/report",
        json={
            "format": "docx",
            "code": CALCULATE_DISCOUNT,
            "params": {"price": 20, "quantity": 6},
            "name": "Calculate Discount",
            "steps": steps,
        },
    )
    assert response.status_code == 200
    assert "wordprocessingml.document" in response.headers["content-type"]

    doc = Document(io.BytesIO(response.content))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Calculate Discount" in all_text
    assert any("108.0" in cell.text for table in doc.tables for row in table.rows for cell in row.cells)


def test_report_endpoint_pdf_is_a_real_pdf():
    steps = _run_debug()
    response = client.post(
        "/debug/report",
        json={"format": "pdf", "code": CALCULATE_DISCOUNT, "name": "Calculate Discount", "steps": steps},
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-")
    assert len(response.content) > 1000


def test_report_endpoint_embeds_client_rasterized_flowchart_image():
    steps = _run_debug()
    response = client.post(
        "/debug/report",
        json={
            "format": "pdf",
            "code": CALCULATE_DISCOUNT,
            "name": "Calculate Discount",
            "steps": steps,
            "flowchartImage": _tiny_png_data_url(),
        },
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")


def test_report_endpoint_ignores_malformed_flowchart_image_instead_of_failing():
    steps = _run_debug()
    response = client.post(
        "/debug/report",
        json={
            "format": "txt",
            "code": CALCULATE_DISCOUNT,
            "name": "Calculate Discount",
            "steps": steps,
            "flowchartImage": "not-a-real-data-url",
        },
    )
    assert response.status_code == 200
    assert "No control-flow diagram image was captured" in response.text


def test_report_endpoint_is_honest_about_missing_params():
    steps = _run_debug(params={})
    response = client.post(
        "/debug/report",
        json={"format": "txt", "code": "DECLARE x NUMBER DEFAULT 1;", "params": {}, "steps": steps},
    )
    assert response.status_code == 200
    assert "does not currently provide a UI" in response.text


# -- Schema-drift regression coverage --------------------------------------
# Every test above only ever exercises CALCULATE_DISCOUNT -- a plain
# DECLARE/SET/IF trace with no `call`, `table`, `cursor`, `branch` (CASE),
# or `loop` (LOOP) DebugStep field ever present. That left a real gap: if
# a future interpreter change renamed/restructured one of those fields,
# nothing here would catch it before it silently broke Download for any
# procedure that actually uses CALL/cursors/CASE/LOOP/tables (see
# CLAUDE.md's Download entry and docs/schema.md for the full DebugStep
# shape). Each test below drives one of those field shapes through the
# REAL /debug -> /debug/report pipeline end to end, the same way the
# frontend's handleDownloadReport does, and asserts on real computed
# values so a schema rename (which would make that value silently vanish
# or KeyError) fails loudly instead of a smaller test drifting unnoticed.

ORDER_TOTAL = """\
CREATE PROCEDURE ComputeSubtotal(IN price NUMBER, IN quantity NUMBER, OUT subtotal NUMBER)
BEGIN
    SET subtotal = price * quantity;
END;

CREATE PROCEDURE OrderTotal()
BEGIN
    DECLARE price NUMBER DEFAULT 20;
    DECLARE quantity NUMBER DEFAULT 6;
    DECLARE taxRate NUMBER DEFAULT 0.08;
    DECLARE subtotal NUMBER DEFAULT 0;
    DECLARE grandTotal NUMBER DEFAULT 0;
    CALL ComputeSubtotal(price, quantity, subtotal);
    SET grandTotal = subtotal + subtotal * taxRate;
END
"""

MANAGE_INVENTORY = """\
CREATE PROCEDURE ManageInventory()
BEGIN
    CREATE TABLE inventory (id NUMBER PRIMARY KEY, item TEXT NOT NULL, qty NUMBER, price NUMBER);
    INSERT INTO inventory VALUES (1, 'Widget', 8, 10);
    INSERT INTO inventory VALUES (2, 'Gadget', 0, 25);
    INSERT INTO inventory VALUES (3, 'Gizmo', 15, 15);
    UPDATE inventory SET qty = qty + 5 WHERE qty < 10;
    UPDATE inventory SET price = price - 2 WHERE price > 12;
    DELETE FROM inventory WHERE id = 2;
END
"""

PRODUCT_PRICE_TOTAL = """\
CREATE PROCEDURE ProductPriceTotal()
BEGIN
    DECLARE total NUMBER DEFAULT 0;
    DECLARE item_name STRING DEFAULT '';
    DECLARE item_price NUMBER DEFAULT 0;
    DECLARE prod_cursor CURSOR FOR SELECT name, price FROM products;
    OPEN prod_cursor;
    WHILE prod_cursor%FOUND DO
        FETCH prod_cursor INTO item_name, item_price;
        SET total = total + item_price;
    END WHILE;
    CLOSE prod_cursor;
END
"""

CLASSIFY_ORDER = """\
CREATE PROCEDURE ClassifyOrder()
BEGIN
    DECLARE quantity NUMBER DEFAULT 12;
    DECLARE unitPrice NUMBER DEFAULT 15;
    DECLARE tier NUMBER DEFAULT 2;
    DECLARE discountRate NUMBER DEFAULT 0;
    DECLARE total NUMBER DEFAULT 0;
    DECLARE sizeLabel NUMBER DEFAULT 0;
    CASE tier
        WHEN 1 THEN
            SET discountRate = 0.05;
        WHEN 2 THEN
            SET discountRate = 0.1;
        WHEN 3 THEN
            SET discountRate = 0.15;
        ELSE
            SET discountRate = 0;
    END CASE;
    SET total = (quantity * unitPrice) * (1 - discountRate);
    CASE
        WHEN total > 200 THEN
            SET sizeLabel = 30;
        WHEN total > 100 THEN
            SET sizeLabel = 20;
        ELSE
            SET sizeLabel = 10;
    END CASE;
END
"""


def _debug_and_report(code, name, fmt="txt"):
    debug_response = client.post("/debug", json={"code": code, "params": {}, "name": name})
    assert debug_response.status_code == 200
    steps = debug_response.json()["steps"]
    report_response = client.post(
        "/debug/report", json={"format": fmt, "code": code, "name": name, "steps": steps}
    )
    assert report_response.status_code == 200
    return report_response


def test_report_endpoint_handles_a_call_based_trace_across_all_formats():
    """OrderTotal's steps include the `call` field (present on every step
    inside the CALLed ComputeSubtotal, per docs/schema.md) -- confirms
    that field's mere presence doesn't trip up report building, and that
    the real cross-procedure computed value still comes through."""
    for fmt in ("txt", "docx", "pdf"):
        response = _debug_and_report(ORDER_TOTAL, "OrderTotal", fmt)
        if fmt == "txt":
            # subtotal = 20*6 = 120; grandTotal = 120 + 120*0.08 = 129.6
            assert "129.6" in response.text
        elif fmt == "docx":
            from docx import Document

            doc = Document(io.BytesIO(response.content))
            all_text = "\n".join(p.text for p in doc.paragraphs)
            cell_text = [c.text for t in doc.tables for r in t.rows for c in r.cells]
            assert "129.6" in all_text or any("129.6" in c for c in cell_text)
        else:
            assert response.content.startswith(b"%PDF-")


def test_report_endpoint_handles_a_table_based_trace():
    """ManageInventory's steps include the `table` field (CREATE TABLE/
    INSERT/UPDATE/DELETE, per docs/schema.md) -- confirms that field's
    presence doesn't break report building even though report.py never
    reads it, and that the real statement types/row data still appear."""
    response = _debug_and_report(MANAGE_INVENTORY, "ManageInventory", "txt")
    text = response.text
    assert "InsertStatement" in text
    assert "UpdateStatement" in text
    assert "DeleteStatement" in text
    assert "Widget" in text


def test_report_endpoint_handles_a_cursor_based_trace():
    """ProductPriceTotal's steps include the `cursor` field -- confirms
    _step_details' cursor formatting still matches the real cursor dict
    shape, and the real summed total (10+25+15=50) comes through."""
    response = _debug_and_report(PRODUCT_PRICE_TOTAL, "ProductPriceTotal", "txt")
    text = response.text
    assert "Cursor prod_cursor" in text
    assert "50" in text


def test_report_endpoint_handles_a_case_statement_trace():
    """ClassifyOrder's steps include the `branch` field on its CASE
    statements (the same field IfStatement uses, per docs/schema.md) --
    confirms _step_details' branch formatting still matches, and the
    real discount/size classification values come through. total =
    (12*15)*(1-0.1) = 162 -> sizeLabel = 20 (100 < 162 <= 200)."""
    response = _debug_and_report(CLASSIFY_ORDER, "ClassifyOrder", "txt")
    text = response.text
    assert "Branch:" in text
    assert "162" in text


def test_report_endpoint_large_trace_stays_fast():
    """Regression test for a real performance bug found while
    investigating a "Download stopped working" report: DOCX/PDF table
    construction (python-docx/reportlab both pay a high per-row cost)
    scaled so badly with step count that even a few-thousand-step trace
    took 10-90+ seconds to download -- indistinguishable from "broken" to
    a user watching a "Preparing..." spinner. Fixed by capping any
    DOCX/PDF table at report._MAX_TABLE_ROWS_FOR_DOCUMENT_EXPORTS rows
    (Text stays untruncated). A 1,000-iteration WHILE loop (3,004 steps)
    used to take ~12s for DOCX/PDF alone; this asserts a generous ceiling
    well above the ~1s it takes now, so ordinary machine variance can't
    make this flaky, but a real regression back to O(n)-per-row-cost
    table building still fails it loudly."""
    import time

    code = """\
CREATE PROCEDURE BigLoop()
BEGIN
    DECLARE i NUMBER DEFAULT 0;
    DECLARE total NUMBER DEFAULT 0;
    WHILE i < 1000 DO
        SET total = total + i;
        SET i = i + 1;
    END WHILE;
END
"""
    debug_response = client.post("/debug", json={"code": code, "params": {}, "name": "BigLoop"})
    assert debug_response.status_code == 200
    steps = debug_response.json()["steps"]
    assert len(steps) > 3000  # confirms this is a genuinely large trace, not a trivial one

    for fmt in ("docx", "pdf"):
        t0 = time.time()
        response = client.post(
            "/debug/report", json={"format": fmt, "code": code, "name": "BigLoop", "steps": steps}
        )
        elapsed = time.time() - t0
        assert response.status_code == 200
        assert elapsed < 5.0, f"{fmt} report took {elapsed:.1f}s for a {len(steps)}-step trace (expected well under 5s)"


def test_report_endpoint_does_not_require_a_fresh_debug_run(monkeypatch):
    """The whole point of this endpoint: it must work from a trace the
    caller already has, without this module re-invoking the interpreter
    itself. Patch app.interpreter.run to explode if it's ever called from
    this endpoint's code path."""
    import app.main as main_module

    steps = _run_debug()  # get a real trace via the *unpatched* /debug first

    def _boom(*args, **kwargs):
        raise AssertionError("report endpoint must not re-run the interpreter")

    monkeypatch.setattr(main_module, "run", _boom)

    response = client.post(
        "/debug/report",
        json={"format": "txt", "code": CALCULATE_DISCOUNT, "name": "Calculate Discount", "steps": steps},
    )
    assert response.status_code == 200
