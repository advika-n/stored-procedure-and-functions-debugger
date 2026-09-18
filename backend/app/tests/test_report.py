import io

from app import report
from app.interpreter import run
from app.parser import parse
from app.tokenizer import tokenize

CALCULATE_DISCOUNT = """\
DECLARE total NUMBER DEFAULT 0;
DECLARE discount NUMBER DEFAULT 0;
SET total = price * quantity;
IF total > 100 THEN
    SET discount = total * 0.1;
ELSE
    SET discount = total * 0.05;
END IF;
SET total = total - discount;
"""


def _real_steps(code=CALCULATE_DISCOUNT, params=None):
    """A real DebugStep trace, produced by actually running the
    interpreter -- never hand-written/fabricated -- so every report test
    below is checking against genuine execution data."""
    ast = parse(tokenize(code))
    steps = run(ast, params or {"price": 20, "quantity": 6})
    return ast, [step.to_dict() for step in steps]


def test_build_report_user_inputs_includes_real_code_and_honest_empty_params():
    _, steps = _real_steps(params={})
    meta, sections = report.build_report(
        code=CALCULATE_DISCOUNT, params={}, procedure_name="Calculate Discount", steps=steps, flowchart_png_bytes=None
    )
    user_inputs = next(s for s in sections if s.title == "User Inputs")
    code_block = next(b for b in user_inputs.blocks if b.kind == "code")
    assert code_block.text == CALCULATE_DISCOUNT
    # No params were sent -- must say so honestly, never fabricate a value.
    joined = " ".join(b.text for b in user_inputs.blocks if b.kind == "paragraph" and b.text)
    assert "none" in joined.lower()
    assert "does not currently provide a UI" in joined


def test_build_report_user_inputs_lists_real_params_when_present():
    _, steps = _real_steps(params={"price": 20, "quantity": 6})
    meta, sections = report.build_report(
        code=CALCULATE_DISCOUNT, params={"price": 20, "quantity": 6}, procedure_name="x", steps=steps, flowchart_png_bytes=None
    )
    user_inputs = next(s for s in sections if s.title == "User Inputs")
    table = next(b for b in user_inputs.blocks if b.kind == "table")
    assert ["price", "20"] in table.rows
    assert ["quantity", "6"] in table.rows


def test_build_report_processing_steps_matches_real_trace():
    _, steps = _real_steps()
    meta, sections = report.build_report(
        code=CALCULATE_DISCOUNT, params={}, procedure_name="x", steps=steps, flowchart_png_bytes=None
    )
    assert meta["step_count"] == len(steps) == 6
    processing = next(s for s in sections if s.title == "Processing Steps")
    table = next(b for b in processing.blocks if b.kind == "table")
    assert len(table.rows) == len(steps)
    # Real line numbers/statement text from the actual steps, not placeholders.
    assert table.rows[0][1] == str(steps[0]["line"])
    assert table.rows[0][3] == steps[0]["statementText"]


def _step_with_error(handler):
    """A minimal, hand-built DebugStep carrying just enough for
    `_step_details` to render its `error` field -- the handled-vs-
    unhandled distinction below doesn't need a real interpreter run, just
    the exact `error` dict shape `docs/schema.md` documents."""
    return {
        "stepNumber": 1,
        "line": 5,
        "nodeType": "FetchCursorNode",
        "statementText": "FETCH cur INTO x;",
        "variables": {},
        "error": {"condition": "NOT_FOUND", "message": "Cursor 'cur' has no more rows to fetch", "handler": handler},
    }


def test_step_details_labels_a_caught_error_as_handled_not_error():
    """Regression test for a real bug: a step caught by a declared
    CONTINUE HANDLER used to render as "ERROR [...] ... caught by X
    handler" in the Processing Steps Details column -- reading as a
    failure when it's actually expected, handled behavior. Confirms the
    fix ("Handled: ..." wording) and that the literal word "ERROR" never
    appears for a caught condition."""
    meta, sections = report.build_report(
        code="x", params={}, procedure_name="x", steps=[_step_with_error("NOT_FOUND handler")], flowchart_png_bytes=None
    )
    details = next(s for s in sections if s.title == "Processing Steps").blocks[1].rows[0][4]
    assert "Handled:" in details
    assert "caught by NOT_FOUND handler" in details
    assert "ERROR" not in details


def test_step_details_no_handler_declared_condition_is_not_labeled_error_either():
    """The other half of the same bug: NOT_FOUND with no CONTINUE HANDLER
    declared at all is STILL non-fatal by the interpreter's own design
    (see interpreter.py's module docstring) -- /debug (and therefore this
    endpoint) only ever sees steps from a run that completed, so a per-step
    `error` here is never the "genuinely unhandled/fatal, halted execution"
    case either; confirms this doesn't get the "ERROR" label."""
    meta, sections = report.build_report(
        code="x", params={}, procedure_name="x", steps=[_step_with_error("unhandled")], flowchart_png_bytes=None
    )
    details = next(s for s in sections if s.title == "Processing Steps").blocks[1].rows[0][4]
    assert "ERROR" not in details
    assert "no handler declared" in details


def test_build_report_has_no_intermediate_results_section():
    """Regression test: an "Intermediate Results" section (one row per
    variable per step) used to sit between Processing Steps and Final
    Output -- deliberately removed as too granular to be useful in a
    written report (see PROMPT_LOG.md). Confirms it's gone AND that the
    remaining four sections are exactly what's left, in order -- not just
    that the removed title is absent (which a typo could satisfy without
    actually testing anything)."""
    _, steps = _real_steps()
    meta, sections = report.build_report(
        code=CALCULATE_DISCOUNT, params={}, procedure_name="x", steps=steps, flowchart_png_bytes=None
    )
    titles = [s.title for s in sections]
    assert "Intermediate Results" not in titles
    assert titles == ["User Inputs", "Processing Steps", "Final Output", "Graphs, Tables & Figures"]


def test_build_report_final_output_uses_last_steps_final_value():
    _, steps = _real_steps(params={"price": 20, "quantity": 6})
    meta, sections = report.build_report(
        code=CALCULATE_DISCOUNT, params={"price": 20, "quantity": 6}, procedure_name="x", steps=steps, flowchart_png_bytes=None
    )
    final = next(s for s in sections if s.title == "Final Output")
    final_table = next(b for b in final.blocks if b.kind == "table")
    total_row = next(r for r in final_table.rows if r[0] == "total")
    # price*quantity = 120 -> >100 branch -> discount 10% -> total = 108.0, the
    # exact same value test_debug_endpoint.py asserts against the real interpreter.
    assert total_row[1] == "108.0"


def test_final_output_no_handler_declared_does_not_say_execution_stopped():
    """Same underlying bug as the two `_step_details` tests above, in the
    Final Output section's own separate rendering: a run's LAST step
    carrying `error.handler == "unhandled"` used to be described as
    "Execution stopped with an unhandled error" -- but `/debug` never
    returns a `steps` list for a run that actually halted (an
    InterpreterError aborts before producing any steps at all), so this
    only ever fires when a non-fatal, no-handler-declared condition (e.g.
    NOT_FOUND) simply happened to be the run's last recorded step, not
    because anything stopped."""
    step = _step_with_error("unhandled")
    step["variables"] = {"x": {"value": None, "type": "null", "changed": False}}
    meta, sections = report.build_report(code="x", params={}, procedure_name="x", steps=[step], flowchart_png_bytes=None)
    final = next(s for s in sections if s.title == "Final Output")
    text = next(b.text for b in final.blocks if b.kind == "paragraph")
    assert "stopped" not in text.lower()
    assert "non-fatal" in text.lower()


def test_build_report_graphs_section_notes_missing_image_honestly():
    _, steps = _real_steps()
    meta, sections = report.build_report(
        code=CALCULATE_DISCOUNT, params={}, procedure_name="x", steps=steps, flowchart_png_bytes=None
    )
    graphs = next(s for s in sections if s.title == "Graphs, Tables & Figures")
    paragraphs = [b.text for b in graphs.blocks if b.kind == "paragraph"]
    assert any("No control-flow diagram image was captured" in p for p in paragraphs)


def test_build_report_graphs_section_includes_real_image_bytes():
    fake_png = b"\x89PNG\r\n\x1a\nfake-but-real-bytes"
    _, steps = _real_steps()
    meta, sections = report.build_report(
        code=CALCULATE_DISCOUNT, params={}, procedure_name="x", steps=steps, flowchart_png_bytes=fake_png
    )
    graphs = next(s for s in sections if s.title == "Graphs, Tables & Figures")
    image_block = next(b for b in graphs.blocks if b.kind == "image")
    assert image_block.image_bytes == fake_png


def test_decode_flowchart_image_handles_data_url_and_raw_base64():
    import base64

    raw = b"hello world"
    encoded = base64.b64encode(raw).decode()
    assert report.decode_flowchart_image(f"data:image/png;base64,{encoded}") == raw
    assert report.decode_flowchart_image(encoded) == raw


def test_decode_flowchart_image_returns_none_for_garbage_or_missing():
    assert report.decode_flowchart_image(None) is None
    assert report.decode_flowchart_image("") is None
    assert report.decode_flowchart_image("not valid base64!!!") is None


def test_generate_report_text_contains_real_data_not_placeholders():
    _, steps = _real_steps(params={"price": 20, "quantity": 6})
    content, content_type, filename = report.generate_report(
        fmt="txt",
        code=CALCULATE_DISCOUNT,
        params={"price": 20, "quantity": 6},
        procedure_name="Calculate Discount",
        steps=steps,
        flowchart_png_bytes=None,
    )
    text = content.decode("utf-8")
    assert content_type == "text/plain; charset=utf-8"
    assert filename.endswith(".txt")
    assert "lorem ipsum" not in text.lower()
    assert "Calculate Discount" in text
    assert "SET total = total - discount" in text
    assert "108.0" in text


def test_generate_report_docx_round_trips_via_python_docx():
    from docx import Document

    _, steps = _real_steps(params={"price": 20, "quantity": 6})
    content, content_type, filename = report.generate_report(
        fmt="docx",
        code=CALCULATE_DISCOUNT,
        params={"price": 20, "quantity": 6},
        procedure_name="Calculate Discount",
        steps=steps,
        flowchart_png_bytes=None,
    )
    assert content_type.endswith("wordprocessingml.document")
    assert filename.endswith(".docx")

    doc = Document(io.BytesIO(content))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Calculate Discount" in all_text
    # At least the Processing Steps table should have real rows -> real cells.
    tables = doc.tables
    assert len(tables) >= 3
    step_table_cells = [cell.text for table in tables for row in table.rows for cell in row.cells]
    assert "IfStatement" in step_table_cells


def test_generate_report_pdf_produces_a_valid_pdf_with_real_content():
    _, steps = _real_steps(params={"price": 20, "quantity": 6})
    content, content_type, filename = report.generate_report(
        fmt="pdf",
        code=CALCULATE_DISCOUNT,
        params={"price": 20, "quantity": 6},
        procedure_name="Calculate Discount",
        steps=steps,
        flowchart_png_bytes=None,
    )
    assert content_type == "application/pdf"
    assert filename.endswith(".pdf")
    assert content.startswith(b"%PDF-")
    assert content.rstrip().endswith(b"%%EOF")
    assert len(content) > 1000  # a real multi-section report, not an empty shell


def test_generate_report_pdf_embeds_a_real_flowchart_image():
    from PIL import Image as PILImage

    # A real (tiny) PNG, not fabricated bytes -- generated with Pillow so
    # reportlab's ImageReader can actually read its real width/height.
    buf = io.BytesIO()
    PILImage.new("RGB", (40, 20), color=(255, 0, 0)).save(buf, format="PNG")
    png_bytes = buf.getvalue()

    _, steps = _real_steps()
    content, _, _ = report.generate_report(
        fmt="pdf",
        code=CALCULATE_DISCOUNT,
        params={},
        procedure_name="Calculate Discount",
        steps=steps,
        flowchart_png_bytes=png_bytes,
    )
    assert content.startswith(b"%PDF-")
    # The raw PNG bytes get re-encoded into the PDF's internal image
    # stream, so we can't substring-match them directly -- a large PDF
    # (bigger than the no-image case) is the practical signal the image
    # was actually embedded rather than skipped.
    assert len(content) > 2000


def test_generate_report_pdf_scales_a_tall_narrow_image_to_fit_the_page():
    """Regression test for a real bug, found live: the image block used
    to scale ONLY by width (`content_width / image_width`), so a tall,
    narrow image -- exactly what a long, vertically-stacked flowchart
    rasterizes to -- could still come out taller than an entire fresh
    page's own frame height after that scaling. reportlab refuses to lay
    out a flowable that doesn't fit at all rather than shrink or split it
    (a "Flowable ... too large ... in frame" exception), which failed the
    WHOLE PDF download with a 500, not just a missing image. Uses a real
    Pillow-generated PNG at a 20:1 aspect ratio -- much taller, relative
    to its width, than LETTER's own content area -- to reproduce the
    original failure if the height-based scale factor were ever removed
    again."""
    from PIL import Image as PILImage

    buf = io.BytesIO()
    PILImage.new("RGB", (200, 4000), color=(20, 30, 60)).save(buf, format="PNG")
    tall_png_bytes = buf.getvalue()

    _, steps = _real_steps()
    content, _, _ = report.generate_report(
        fmt="pdf",
        code=CALCULATE_DISCOUNT,
        params={},
        procedure_name="Calculate Discount",
        steps=steps,
        flowchart_png_bytes=tall_png_bytes,
    )
    assert content.startswith(b"%PDF-")
    assert content.rstrip().endswith(b"%%EOF")


def test_cap_table_rows_leaves_small_tables_untouched():
    rows = [["a", "1"], ["b", "2"]]
    capped, note = report._cap_table_rows(rows, max_rows=500)
    assert capped == rows
    assert note is None


def test_cap_table_rows_truncates_and_notes_large_tables():
    rows = [[str(i)] for i in range(600)]
    capped, note = report._cap_table_rows(rows, max_rows=500)
    assert len(capped) == 500
    assert capped == rows[:500]
    assert note is not None
    assert "100" in note  # 600 - 500 omitted
    assert "Text export" in note


def test_generate_report_docx_caps_large_tables_but_txt_stays_complete():
    """Regression test for the report-download performance bug (see
    test_report_endpoint.py's test_report_endpoint_large_trace_stays_fast
    for the full story): DOCX/PDF must cap any table that scales with
    step count, but the plain-text renderer -- which has no comparable
    per-row cost -- must keep the complete, untruncated data."""
    from docx import Document

    big_step_count = report._MAX_TABLE_ROWS_FOR_DOCUMENT_EXPORTS + 50
    steps = [
        {
            "stepNumber": i + 1,
            "line": i + 1,
            "nodeType": "SetStatement",
            "statementText": f"SET x = {i};",
            "variables": {},
        }
        for i in range(big_step_count)
    ]

    txt_content, _, _ = report.generate_report(
        fmt="txt", code="x", params={}, procedure_name="Big", steps=steps, flowchart_png_bytes=None
    )
    text = txt_content.decode("utf-8")
    assert f"SET x = {big_step_count - 1};" in text  # the real last step, present -- not truncated
    assert "omitted" not in text

    docx_content, _, _ = report.generate_report(
        fmt="docx", code="x", params={}, procedure_name="Big", steps=steps, flowchart_png_bytes=None
    )
    doc = Document(io.BytesIO(docx_content))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "omitted" in all_text
    assert "Text export" in all_text


def test_generate_report_rejects_unsupported_format():
    import pytest

    _, steps = _real_steps()
    with pytest.raises(ValueError):
        report.generate_report(
            fmt="exe", code=CALCULATE_DISCOUNT, params={}, procedure_name="x", steps=steps, flowchart_png_bytes=None
        )
