"""Multi-format (PDF / DOCX / plain text) report export for a completed
debug run.

Design: build ONE format-agnostic intermediate representation (a list of
`ReportSection`, each holding `ReportBlock`s) from the exact same data
every successful `/debug` response already carries -- the source code,
the params that were sent, and the `DebugStep` trace (see `interpreter.py`
and `CLAUDE.md` §4 for that shape) -- then feed that single IR into three
independent renderers, one per format. This keeps the three formats'
*content* impossible to drift apart: a change to what a report says only
ever happens once, in `build_report()`. Each renderer only has to know
how to lay out paragraphs/tables/images in its own library.

Library choices (see HANDOFF.md for the "why" in one place too):
  - PDF: **reportlab**, not weasyprint. weasyprint needs the GTK/Pango/
    Cairo native libraries on Windows, which aren't guaranteed present
    and are painful to install reliably; reportlab is pure Python, pip-
    installs cleanly on any platform, and this project's dev environment
    is Windows -- "simpler to integrate with the existing project
    structure" pointed at reportlab, not weasyprint's HTML-to-PDF route.
  - DOCX: **python-docx** -- the standard, only-real-option library for
    writing .docx from Python.
  - The flowchart image is rasterized to a PNG **client-side** (see
    `frontend/src/cfg.js`'s `rasterizeSvgToPng`) from the Mermaid SVG the
    Debugger has already rendered, and sent to this module as PNG bytes.
    Converting Mermaid's SVG output (which relies on `<foreignObject>`
    HTML labels) to PDF/DOCX-embeddable images purely server-side in
    Python has no reliable, dependency-light option; a browser can
    rasterize its own already-rendered SVG through a `<canvas>` in a few
    lines, so that's where this happens instead.
"""

from __future__ import annotations

import base64
import binascii
import io
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone
from xml.sax.saxutils import escape as _xml_escape


# -- intermediate representation ---------------------------------------------


@dataclass
class ReportBlock:
    kind: str  # "paragraph" | "code" | "table" | "image"
    text: str | None = None
    headers: list[str] | None = None
    rows: list[list[str]] | None = None
    col_ratios: list[float] | None = None  # PDF column-width hints only
    image_bytes: bytes | None = None
    caption: str | None = None


@dataclass
class ReportSection:
    title: str
    blocks: list[ReportBlock] = field(default_factory=list)


def _fmt_value(value) -> str:
    """Render a DebugStep variable/return value for report display --
    mirrors the frontend's formatValue() (DebuggerPage.jsx): an em-dash
    placeholder for null/undefined (never Python's "None"), and quoted
    strings so an empty string ("") reads as a real, present value
    rather than looking identical to "no value" in a plain-text table."""
    if value is None:
        return "—"
    if isinstance(value, str):
        return f'"{value}"'
    return str(value)


def _step_details(step: dict) -> str:
    """One line summarizing whatever this step's branch/loop/cursor/error
    fields say happened -- kept single-line (joined with "; ") on purpose
    so the plain-text renderer's fixed-width table columns stay aligned;
    the PDF/DOCX renderers render it as ordinary wrapped text either way."""
    parts = []
    branch = step.get("branch")
    if branch:
        parts.append(
            f"Branch: {branch['condition']} -> {'TRUE' if branch['result'] else 'FALSE'} ({branch['path']} taken)"
        )
    loop = step.get("loop")
    if loop:
        parts.append(
            f"Loop iteration {loop['iteration']}: {loop['condition']} -> {'TRUE' if loop['result'] else 'FALSE'}"
        )
    cursor = step.get("cursor")
    if cursor:
        current_row = cursor.get("currentRow") or {}
        row_desc = ", ".join(f"{k}={v}" for k, v in current_row.items()) if current_row else "—"
        parts.append(
            f"Cursor {cursor['name']}: row {cursor['rowIndex']} "
            f"({'more rows' if cursor['hasMore'] else 'no more rows'}), current row: {row_desc}"
        )
    error = step.get("error")
    if error:
        caught = "unhandled" if error["handler"] == "unhandled" else f"caught by {error['handler']}"
        parts.append(f"ERROR [{error['condition']}] {error['message']} ({caught})")
    return "; ".join(parts) if parts else "—"


def build_report(
    *,
    code: str,
    params: dict,
    procedure_name: str,
    steps: list[dict],
    flowchart_png_bytes: bytes | None,
) -> tuple[dict, list[ReportSection]]:
    """Build the format-agnostic report content. `steps` is the exact
    camelCase DebugStep list a `/debug` response already returns -- this
    function does no re-execution, honoring the "reuse the session's
    data" requirement rather than needing the code re-run server-side."""
    meta = {
        "title": procedure_name or "Custom procedure",
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"),
        "step_count": len(steps),
    }

    sections: list[ReportSection] = []

    # -- a. User Inputs -------------------------------------------------
    input_blocks = [
        ReportBlock(kind="paragraph", text=f"Procedure/Function: {meta['title']}"),
        ReportBlock(kind="paragraph", text="SQL source, exactly as entered:"),
        ReportBlock(kind="code", text=code),
    ]
    if params:
        input_blocks.append(ReportBlock(kind="paragraph", text="Parameters supplied for this run:"))
        input_blocks.append(
            ReportBlock(
                kind="table",
                headers=["Name", "Value"],
                rows=[[name, _fmt_value(value)] for name, value in params.items()],
                col_ratios=[0.35, 0.65],
            )
        )
    else:
        # Honest, not fabricated: the frontend has no UI yet for entering
        # external parameter values before running (see the Help tab and
        # DebuggerPage.jsx's handleDebug, which always sends params={}).
        # Printing a fake example value here would misrepresent this run.
        input_blocks.append(
            ReportBlock(
                kind="paragraph",
                text=(
                    "Parameters supplied for this run: none. This application does not currently "
                    "provide a UI for entering external parameter values before running a procedure "
                    "or function -- every run must be self-contained via DECLARE ... DEFAULT inside "
                    "the body itself (see the Help tab's \"What inputs you can give it\" section)."
                ),
            )
        )
    sections.append(ReportSection("User Inputs", input_blocks))

    # -- b. Processing Steps ---------------------------------------------
    step_rows = [
        [
            str(step["stepNumber"]),
            str(step["line"]),
            step["nodeType"],
            step["statementText"],
            _step_details(step),
        ]
        for step in steps
    ]
    sections.append(
        ReportSection(
            "Processing Steps",
            [
                ReportBlock(kind="paragraph", text=f"{len(steps)} step(s) executed, in order:"),
                ReportBlock(
                    kind="table",
                    headers=["Step", "Line", "Type", "Statement", "Details"],
                    rows=step_rows,
                    col_ratios=[0.06, 0.06, 0.14, 0.34, 0.40],
                ),
            ],
        )
    )

    # -- c. Intermediate Results (variable state at every step) ---------
    var_rows = []
    for step in steps:
        for name, entry in (step.get("variables") or {}).items():
            var_rows.append(
                [
                    str(step["stepNumber"]),
                    name,
                    _fmt_value(entry["value"]),
                    entry["type"],
                    "yes" if entry["changed"] else "",
                    "yes" if entry.get("isOutput") else "",
                ]
            )
    sections.append(
        ReportSection(
            "Intermediate Results",
            [
                ReportBlock(
                    kind="paragraph",
                    text=(
                        'Variable state captured at every step ("Changed" marks the step that set that '
                        'exact value; "Output" marks an OUT/INOUT parameter):'
                    ),
                ),
                ReportBlock(
                    kind="table",
                    headers=["Step", "Variable", "Value", "Type", "Changed", "Output"],
                    rows=var_rows,
                    col_ratios=[0.08, 0.20, 0.24, 0.16, 0.16, 0.16],
                ),
            ],
        )
    )

    # -- d. Final Output ---------------------------------------------------
    last_step = steps[-1] if steps else None
    final_blocks: list[ReportBlock] = []
    if last_step is None:
        final_blocks.append(ReportBlock(kind="paragraph", text="No steps were recorded for this run."))
    else:
        error = last_step.get("error")
        return_value = last_step.get("returnValue")
        if return_value is not None:
            final_blocks.append(
                ReportBlock(
                    kind="paragraph",
                    text=f"Function returned: {_fmt_value(return_value['value'])} ({return_value['type']})",
                )
            )
        elif error is not None and error["handler"] == "unhandled":
            final_blocks.append(
                ReportBlock(
                    kind="paragraph",
                    text=(
                        f"Execution stopped with an unhandled error at step {last_step['stepNumber']} "
                        f"(line {last_step['line']}): [{error['condition']}] {error['message']}"
                    ),
                )
            )
        else:
            final_blocks.append(
                ReportBlock(
                    kind="paragraph",
                    text=f"Execution completed normally after {len(steps)} step(s), with no unhandled error.",
                )
            )
        final_vars = last_step.get("variables") or {}
        if final_vars:
            final_blocks.append(ReportBlock(kind="paragraph", text="Final variable state:"))
            final_blocks.append(
                ReportBlock(
                    kind="table",
                    headers=["Variable", "Value", "Type", "Output"],
                    rows=[
                        [name, _fmt_value(entry["value"]), entry["type"], "yes" if entry.get("isOutput") else ""]
                        for name, entry in final_vars.items()
                    ],
                    col_ratios=[0.28, 0.32, 0.20, 0.20],
                )
            )
    sections.append(ReportSection("Final Output", final_blocks))

    # -- e. Graphs, Tables & Figures -------------------------------------
    graph_blocks: list[ReportBlock] = []
    if flowchart_png_bytes:
        graph_blocks.append(
            ReportBlock(
                kind="image",
                image_bytes=flowchart_png_bytes,
                caption="Control-flow diagram, as rendered in the Debugger.",
            )
        )
    else:
        graph_blocks.append(
            ReportBlock(
                kind="paragraph",
                text=(
                    "No control-flow diagram image was captured for this export "
                    "(plain-text exports never include one; a PDF/Document export only omits "
                    "it if the flowchart hadn't finished rendering in the Debugger yet)."
                ),
            )
        )
    type_counts = Counter(step["nodeType"] for step in steps)
    if type_counts:
        graph_blocks.append(ReportBlock(kind="paragraph", text="Steps by statement type:"))
        graph_blocks.append(
            ReportBlock(
                kind="table",
                headers=["Statement type", "Step count"],
                rows=[[node_type, str(count)] for node_type, count in type_counts.most_common()],
                col_ratios=[0.6, 0.4],
            )
        )
    sections.append(ReportSection("Graphs, Tables & Figures", graph_blocks))

    return meta, sections


def decode_flowchart_image(data: str | None) -> bytes | None:
    """Best-effort decode of the client-sent PNG (a data URL or raw
    base64). The diagram is optional content ("where applicable" per the
    spec) -- malformed image data degrades to "no diagram" text rather
    than failing the whole report."""
    if not data:
        return None
    if "," in data and data.strip().lower().startswith("data:"):
        data = data.split(",", 1)[1]
    try:
        return base64.b64decode(data, validate=True)
    except (binascii.Error, ValueError):
        return None


# -- plain text renderer ------------------------------------------------------


def _render_text_table(headers: list[str], rows: list[list[str]]) -> list[str]:
    if not rows:
        return ["(none)"]
    widths = [len(h) for h in headers]
    str_rows = [[str(cell) for cell in row] for row in rows]
    for row in str_rows:
        for i, cell in enumerate(row):
            widths[i] = max(widths[i], len(cell))

    def fmt_row(cells: list[str]) -> str:
        return " | ".join(cell.ljust(widths[i]) for i, cell in enumerate(cells))

    lines = [fmt_row(headers), "-+-".join("-" * w for w in widths)]
    lines.extend(fmt_row(row) for row in str_rows)
    return lines


def render_text(meta: dict, sections: list[ReportSection]) -> bytes:
    lines = [
        "=" * 78,
        "STORED PROCEDURE / FUNCTION DEBUG REPORT",
        f"Procedure: {meta['title']}",
        f"Generated: {meta['generated_at']}",
        f"Steps executed: {meta['step_count']}",
        "=" * 78,
    ]
    for section in sections:
        lines.append("")
        lines.append(section.title.upper())
        lines.append("-" * len(section.title))
        for block in section.blocks:
            if block.kind == "paragraph":
                lines.append(block.text or "")
            elif block.kind == "code":
                lines.append("")
                lines.extend(f"    {line}" for line in (block.text or "").splitlines())
                lines.append("")
            elif block.kind == "table":
                lines.append("")
                lines.extend(_render_text_table(block.headers or [], block.rows or []))
                lines.append("")
            elif block.kind == "image":
                lines.append(
                    f"[Figure: {block.caption or 'image'} -- see the PDF or Document export to view it; "
                    "plain text cannot embed images.]"
                )
    return ("\n".join(lines) + "\n").encode("utf-8")


# -- DOCX renderer -------------------------------------------------------------


def render_docx(meta: dict, sections: list[ReportSection]) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    doc.add_heading("Stored Procedure / Function Debug Report", level=0)
    doc.add_paragraph().add_run(f"Procedure: {meta['title']}").bold = True
    doc.add_paragraph(f"Generated: {meta['generated_at']}")
    doc.add_paragraph(f"Steps executed: {meta['step_count']}")

    for section in sections:
        doc.add_heading(section.title, level=1)
        for block in section.blocks:
            if block.kind == "paragraph":
                doc.add_paragraph(block.text or "")
            elif block.kind == "code":
                run = doc.add_paragraph().add_run(block.text or "")
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif block.kind == "table":
                _add_docx_table(doc, block.headers or [], block.rows or [])
            elif block.kind == "image":
                if block.caption:
                    caption_run = doc.add_paragraph().add_run(block.caption)
                    caption_run.italic = True
                doc.add_picture(io.BytesIO(block.image_bytes), width=Inches(6))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_docx_table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=max(len(headers), 1))
    table.style = "Table Grid"  # a standard style shipped with python-docx's default template
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)


# -- PDF renderer (reportlab) ---------------------------------------------------


def render_pdf(meta: dict, sections: list[ReportSection]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    left_margin = right_margin = top_margin = bottom_margin = 0.6 * inch
    doc = SimpleDocTemplate(
        buf,
        pagesize=LETTER,
        leftMargin=left_margin,
        rightMargin=right_margin,
        topMargin=top_margin,
        bottomMargin=bottom_margin,
    )
    content_width = LETTER[0] - left_margin - right_margin

    styles = getSampleStyleSheet()
    body_style = styles["Normal"]
    code_style = ParagraphStyle("ReportCode", parent=body_style, fontName="Courier", fontSize=8, leading=10)
    cell_style = ParagraphStyle("ReportCell", parent=body_style, fontSize=7.5, leading=9.5)
    header_cell_style = ParagraphStyle(
        "ReportHeaderCell", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold"
    )

    def esc(text: str) -> str:
        return _xml_escape(str(text)).replace("\n", "<br/>")

    story = [
        Paragraph("Stored Procedure / Function Debug Report", styles["Title"]),
        Paragraph(f"<b>Procedure:</b> {esc(meta['title'])}", body_style),
        Paragraph(f"<b>Generated:</b> {esc(meta['generated_at'])}", body_style),
        Paragraph(f"<b>Steps executed:</b> {meta['step_count']}", body_style),
        Spacer(1, 0.2 * inch),
    ]

    for section in sections:
        story.append(Paragraph(esc(section.title), styles["Heading1"]))
        for block in section.blocks:
            if block.kind == "paragraph":
                story.append(Paragraph(esc(block.text or ""), body_style))
                story.append(Spacer(1, 0.08 * inch))
            elif block.kind == "code":
                for line in (block.text or "").splitlines() or [""]:
                    story.append(Paragraph(esc(line) or "&nbsp;", code_style))
                story.append(Spacer(1, 0.1 * inch))
            elif block.kind == "table":
                story.append(
                    _build_pdf_table(
                        block.headers or [],
                        block.rows or [],
                        content_width,
                        block.col_ratios,
                        cell_style,
                        header_cell_style,
                        esc,
                    )
                )
                story.append(Spacer(1, 0.15 * inch))
            elif block.kind == "image":
                if block.caption:
                    story.append(Paragraph(f"<i>{esc(block.caption)}</i>", body_style))
                image_width, image_height = ImageReader(io.BytesIO(block.image_bytes)).getSize()
                scale = min(1.0, content_width / image_width) if image_width else 1.0
                story.append(Image(io.BytesIO(block.image_bytes), width=image_width * scale, height=image_height * scale))
                story.append(Spacer(1, 0.15 * inch))
        story.append(Spacer(1, 0.1 * inch))

    doc.build(story)
    return buf.getvalue()


def _build_pdf_table(headers, rows, content_width, col_ratios, cell_style, header_cell_style, esc):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    col_count = max(len(headers), 1)
    ratios = col_ratios if col_ratios and len(col_ratios) == col_count else [1 / col_count] * col_count
    col_widths = [content_width * r for r in ratios]

    data = [[Paragraph(esc(h), header_cell_style) for h in headers]]
    if rows:
        for row in rows:
            data.append([Paragraph(esc(cell), cell_style) for cell in row])
    else:
        data.append([Paragraph("(none)", cell_style)] + [Paragraph("", cell_style) for _ in range(col_count - 1)])

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


# -- entry point ---------------------------------------------------------------

_CONTENT_TYPES = {
    "txt": "text/plain; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def generate_report(
    *,
    fmt: str,
    code: str,
    params: dict,
    procedure_name: str,
    steps: list[dict],
    flowchart_png_bytes: bytes | None,
) -> tuple[bytes, str, str]:
    """Returns (file_bytes, content_type, filename)."""
    if fmt not in _CONTENT_TYPES:
        raise ValueError(f"Unsupported report format: {fmt!r}")

    meta, sections = build_report(
        code=code,
        params=params,
        procedure_name=procedure_name,
        steps=steps,
        flowchart_png_bytes=flowchart_png_bytes,
    )

    if fmt == "txt":
        content = render_text(meta, sections)
    elif fmt == "docx":
        content = render_docx(meta, sections)
    else:
        content = render_pdf(meta, sections)

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    safe_name = "".join(c if c.isalnum() or c in "-_ " else "_" for c in (procedure_name or "procedure")).strip() or "procedure"
    filename = f"debug-report-{safe_name}-{timestamp}.{fmt}"
    return content, _CONTENT_TYPES[fmt], filename
